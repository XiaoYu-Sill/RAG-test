from __future__ import annotations

import uuid
from pathlib import Path
from typing import Iterable

import pandas as pd
import textract
from docx import Document as DocxDocument
from langchain.schema import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from pptx import Presentation

SUPPORTED_EXTENSIONS = {".pdf", ".doc", ".docx", ".csv", ".xlsx", ".pptx"}


def load_documents(file_path: Path) -> list[Document]:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(file_path)
    if suffix == ".docx":
        return _load_docx(file_path)
    if suffix == ".doc":
        return _load_doc(file_path)
    if suffix == ".csv":
        return _load_csv(file_path)
    if suffix == ".xlsx":
        return _load_xlsx(file_path)
    if suffix == ".pptx":
        return _load_pptx(file_path)
    raise ValueError(f"Unsupported file type: {suffix}")


def split_documents(documents: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=120)
    chunks = splitter.split_documents(list(documents))
    for index, chunk in enumerate(chunks, start=1):
        chunk.metadata = {**chunk.metadata, "chunk_id": index}
    return chunks


def build_records(documents: Iterable[Document]) -> list[tuple[str, Document]]:
    records: list[tuple[str, Document]] = []
    for doc in documents:
        record_id = str(uuid.uuid4())
        records.append((record_id, doc))
    return records


def _base_metadata(file_path: Path) -> dict[str, str]:
    return {
        "source_name": file_path.name,
        "source_path": str(file_path),
        "source_type": file_path.suffix.lower().lstrip("."),
    }


def _load_pdf(file_path: Path) -> list[Document]:
    reader = PdfReader(str(file_path))
    documents: list[Document] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if not text.strip():
            continue
        metadata = {
            **_base_metadata(file_path),
            "page": index,
            "location": f"第{index}页",
        }
        documents.append(Document(page_content=text, metadata=metadata))
    return documents


def _load_docx(file_path: Path) -> list[Document]:
    docx = DocxDocument(str(file_path))
    documents: list[Document] = []
    for index, paragraph in enumerate(docx.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        metadata = {
            **_base_metadata(file_path),
            "paragraph": index,
            "location": f"第{index}段",
        }
        documents.append(Document(page_content=text, metadata=metadata))
    return documents


def _load_doc(file_path: Path) -> list[Document]:
    text = textract.process(str(file_path)).decode("utf-8", errors="ignore")
    metadata = {**_base_metadata(file_path), "location": "全文"}
    return [Document(page_content=text, metadata=metadata)] if text.strip() else []


def _load_csv(file_path: Path) -> list[Document]:
    df = pd.read_csv(file_path)
    documents: list[Document] = []
    for index, row in df.iterrows():
        row_text = ", ".join(f"{col}: {row[col]}" for col in df.columns)
        metadata = {
            **_base_metadata(file_path),
            "row": int(index) + 1,
            "location": f"第{int(index) + 1}行",
        }
        documents.append(Document(page_content=row_text, metadata=metadata))
    return documents


def _load_xlsx(file_path: Path) -> list[Document]:
    documents: list[Document] = []
    workbook = pd.ExcelFile(file_path)
    for sheet_name in workbook.sheet_names:
        df = workbook.parse(sheet_name)
        for index, row in df.iterrows():
            row_text = ", ".join(f"{col}: {row[col]}" for col in df.columns)
            metadata = {
                **_base_metadata(file_path),
                "sheet": sheet_name,
                "row": int(index) + 1,
                "location": f"{sheet_name} 第{int(index) + 1}行",
            }
            documents.append(Document(page_content=row_text, metadata=metadata))
    return documents


def _load_pptx(file_path: Path) -> list[Document]:
    presentation = Presentation(str(file_path))
    documents: list[Document] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                texts.append(shape.text)
        text = "\n".join(texts).strip()
        if not text:
            continue
        metadata = {
            **_base_metadata(file_path),
            "slide": index,
            "location": f"第{index}页幻灯片",
        }
        documents.append(Document(page_content=text, metadata=metadata))
    return documents
