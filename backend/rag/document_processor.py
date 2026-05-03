import os
import uuid
import logging
from typing import List, Dict, Any

from langchain.text_splitter import RecursiveCharacterTextSplitter

from config import CHUNK_SIZE, CHUNK_OVERLAP

logger = logging.getLogger(__name__)


class DocumentProcessor:
    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=CHUNK_SIZE,
            chunk_overlap=CHUNK_OVERLAP,
            separators=["\n\n", "\n", "。", "！", "？", ".", "!", "?", " ", ""],
        )

    def load_document(self, file_path: str) -> List[Dict[str, Any]]:
        """Load a document and return list of {text, metadata} dicts."""
        ext = os.path.splitext(file_path)[1].lower()
        filename = os.path.basename(file_path)

        if ext == ".pdf":
            return self._load_pdf(file_path, filename)
        elif ext in (".docx", ".doc"):
            return self._load_docx(file_path, filename)
        elif ext == ".csv":
            return self._load_csv(file_path, filename)
        elif ext == ".xlsx":
            return self._load_xlsx(file_path, filename)
        elif ext == ".pptx":
            return self._load_pptx(file_path, filename)
        elif ext in (".txt", ".md"):
            return self._load_text(file_path, filename)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

    def _load_pdf(self, file_path: str, filename: str) -> List[Dict[str, Any]]:
        try:
            import fitz  # PyMuPDF
        except ImportError:
            raise ImportError("PyMuPDF (fitz) is required for PDF processing. Install with: pip install pymupdf")

        pages = []
        doc = fitz.open(file_path)
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append({"text": text, "page": page_num})
        doc.close()
        return self._split_pages(pages, filename)

    def _load_docx(self, file_path: str, filename: str) -> List[Dict[str, Any]]:
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx is required. Install with: pip install python-docx")

        doc = Document(file_path)
        full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        pages = [{"text": full_text, "page": 1}]
        return self._split_pages(pages, filename)

    def _load_csv(self, file_path: str, filename: str) -> List[Dict[str, Any]]:
        try:
            import pandas as pd
        except ImportError:
            raise ImportError("pandas is required. Install with: pip install pandas")

        df = pd.read_csv(file_path, encoding="utf-8-sig")
        text = df.to_string(index=False)
        pages = [{"text": text, "page": 1}]
        return self._split_pages(pages, filename)

    def _load_xlsx(self, file_path: str, filename: str) -> List[Dict[str, Any]]:
        try:
            import openpyxl
        except ImportError:
            raise ImportError("openpyxl is required. Install with: pip install openpyxl")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        pages = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                row_text = "\t".join(str(c) if c is not None else "" for c in row)
                if row_text.strip():
                    rows.append(row_text)
            if rows:
                pages.append({"text": f"[Sheet: {sheet_name}]\n" + "\n".join(rows), "page": sheet_name})
        return self._split_pages(pages, filename)

    def _load_pptx(self, file_path: str, filename: str) -> List[Dict[str, Any]]:
        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError("python-pptx is required. Install with: pip install python-pptx")

        prs = Presentation(file_path)
        pages = []
        for slide_num, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            texts.append(t)
            if texts:
                pages.append({"text": "\n".join(texts), "page": slide_num})
        return self._split_pages(pages, filename)

    def _load_text(self, file_path: str, filename: str) -> List[Dict[str, Any]]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        pages = [{"text": text, "page": 1}]
        return self._split_pages(pages, filename)

    def _split_pages(self, pages: List[Dict[str, Any]], filename: str) -> List[Dict[str, Any]]:
        """Split page texts into chunks, preserving page metadata."""
        chunks = []
        for page_info in pages:
            page_text = page_info["text"]
            page_num = page_info["page"]
            if not page_text.strip():
                continue
            splits = self.text_splitter.split_text(page_text)
            for split_text in splits:
                if split_text.strip():
                    chunks.append({
                        "text": split_text.strip(),
                        "metadata": {
                            "source": filename,
                            "page": page_num,
                            "chunk_id": str(uuid.uuid4()),
                        },
                    })
        logger.info(f"Document '{filename}' split into {len(chunks)} chunks")
        return chunks
